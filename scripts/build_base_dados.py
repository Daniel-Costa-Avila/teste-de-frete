from __future__ import annotations

import argparse
import re
import sqlite3
import os
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from app.infra.product_sheet import parse_products_file  # noqa: E402


@dataclass(frozen=True)
class CepRow:
    uf: str
    cidade: str
    regiao: str
    cep: str


def _normalize_cep(value: object) -> str:
    raw = str(value or "").strip()
    digits = re.sub(r"\D", "", raw)
    if len(digits) != 8:
        return ""
    return f"{digits[:5]}-{digits[5:]}"


def _iter_ceps_from_xlsx(path: Path) -> Iterable[CepRow]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        header_row = None
        idx_uf = idx_cidade = idx_regiao = idx_cep = None

        for i in range(1, 30):
            values = [str(ws.cell(i, j).value or "").strip().lower() for j in range(1, 20)]
            if "cep" in values and "uf" in values:
                header_row = i
                idx_uf = values.index("uf")
                idx_cidade = values.index("cidade") if "cidade" in values else None
                idx_regiao = values.index("regiao") if "regiao" in values else None
                idx_cep = values.index("cep")
                break

        if header_row is None or idx_uf is None or idx_cep is None:
            raise ValueError("Nao foi possivel localizar cabecalhos UF/Cidade/Regiao/CEP na planilha.")

        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            if not row:
                continue

            uf = str(row[idx_uf] or "").strip().upper() if idx_uf < len(row) else ""
            cidade = str(row[idx_cidade] or "").strip() if idx_cidade is not None and idx_cidade < len(row) else ""
            regiao = str(row[idx_regiao] or "").strip() if idx_regiao is not None and idx_regiao < len(row) else ""
            cep = _normalize_cep(row[idx_cep] if idx_cep < len(row) else "")

            if not cep:
                continue

            yield CepRow(uf=uf, cidade=cidade, regiao=regiao, cep=cep)
    finally:
        wb.close()


def _ensure_schema(conn: sqlite3.Connection) -> None:
    conn.executescript(
        """
        PRAGMA foreign_keys = ON;

        CREATE TABLE IF NOT EXISTS ceps (
            cep TEXT PRIMARY KEY,
            uf TEXT NOT NULL,
            cidade TEXT,
            regiao TEXT,
            created_at TEXT NOT NULL DEFAULT (datetime('now'))
        );

        CREATE TABLE IF NOT EXISTS products (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            group_name TEXT,
            product_name TEXT,
            product_id TEXT UNIQUE,
            url TEXT NOT NULL UNIQUE,
            created_at TEXT NOT NULL DEFAULT (datetime('now'))
        );

        CREATE TABLE IF NOT EXISTS product_ceps (
            product_id INTEGER NOT NULL,
            cep TEXT NOT NULL,
            created_at TEXT NOT NULL DEFAULT (datetime('now')),
            PRIMARY KEY (product_id, cep),
            FOREIGN KEY (product_id) REFERENCES products(id) ON DELETE CASCADE,
            FOREIGN KEY (cep) REFERENCES ceps(cep) ON DELETE RESTRICT
        );

        CREATE INDEX IF NOT EXISTS idx_products_group_name ON products(group_name);
        CREATE INDEX IF NOT EXISTS idx_product_ceps_cep ON product_ceps(cep);
        """
    )


def _upsert_ceps(conn: sqlite3.Connection, ceps: Iterable[CepRow]) -> int:
    total = 0
    for row in ceps:
        conn.execute(
            """
            INSERT INTO ceps (cep, uf, cidade, regiao)
            VALUES (?, ?, ?, ?)
            ON CONFLICT(cep) DO UPDATE SET
                uf=excluded.uf,
                cidade=excluded.cidade,
                regiao=excluded.regiao
            """,
            (row.cep, row.uf, row.cidade, row.regiao),
        )
        total += 1
    return total


def _import_products(conn: sqlite3.Connection, products_file: Path) -> tuple[int, int]:
    data = products_file.read_bytes()
    rows = parse_products_file(products_file.name, data)
    product_count = 0
    rel_count = 0

    for row in rows:
        url = (row.url or "").strip()
        if not url:
            continue

        conn.execute(
            """
            INSERT INTO products (group_name, product_name, product_id, url)
            VALUES (?, ?, ?, ?)
            ON CONFLICT(url) DO UPDATE SET
                group_name=excluded.group_name,
                product_name=excluded.product_name,
                product_id=COALESCE(excluded.product_id, products.product_id)
            """,
            (
                (row.group or "").strip() or None,
                (row.product_name or "").strip() or None,
                (row.product_id or "").strip() or None,
                url,
            ),
        )

        db_id = conn.execute("SELECT id FROM products WHERE url = ?", (url,)).fetchone()[0]
        product_count += 1

        for cep in row.ceps:
            if not _normalize_cep(cep):
                continue
            conn.execute(
                "INSERT OR IGNORE INTO product_ceps (product_id, cep) VALUES (?, ?)",
                (db_id, cep),
            )
            rel_count += 1

    return product_count, rel_count


def _default_products_file() -> str:
    env_path = (os.getenv("PRODUCTS_FILE_PATH") or os.getenv("PRODUCTS_BASE_FILE") or "").strip()
    if env_path:
        path = Path(env_path)
        if path.is_absolute():
            if path.exists():
                return str(path)
        else:
            rooted = ROOT / path
            if rooted.exists():
                return str(rooted)

    for candidate in (
        ROOT / "Base de teste.xlsx",
        ROOT / "base_de_teste.xlsx",
        ROOT / "artifacts" / "produtos_entrada_template.xlsx",
    ):
        if candidate.exists():
            return str(candidate)

    return ""


def main() -> None:
    parser = argparse.ArgumentParser(description="Cria base local de produtos e CEPs em SQLite.")
    parser.add_argument("--ceps-xlsx", default="Faixas de CEP.xlsx", help="Planilha de CEPs.")
    parser.add_argument(
        "--products-file",
        default=_default_products_file(),
        help="Planilha/CSV de produtos (opcional). Se omitido, tenta usar 'Base de teste.xlsx'.",
    )
    parser.add_argument("--db", default=str(ROOT / "artifacts" / "base_frete.db"), help="Caminho do banco SQLite.")
    args = parser.parse_args()

    ceps_path = Path(args.ceps_xlsx)
    db_path = Path(args.db)
    products_path = Path(args.products_file) if args.products_file else None

    if not ceps_path.exists():
        raise SystemExit(f"Arquivo de CEPs nao encontrado: {ceps_path}")

    db_path.parent.mkdir(parents=True, exist_ok=True)

    conn = sqlite3.connect(db_path)
    try:
        _ensure_schema(conn)

        cep_rows = list(_iter_ceps_from_xlsx(ceps_path))
        imported_ceps = _upsert_ceps(conn, cep_rows)

        imported_products = 0
        imported_relations = 0
        if products_path and products_path.exists():
            imported_products, imported_relations = _import_products(conn, products_path)

        conn.commit()

        total_ceps = conn.execute("SELECT COUNT(*) FROM ceps").fetchone()[0]
        total_products = conn.execute("SELECT COUNT(*) FROM products").fetchone()[0]
        total_rel = conn.execute("SELECT COUNT(*) FROM product_ceps").fetchone()[0]

        print("Base criada/atualizada com sucesso")
        print(f"DB: {db_path}")
        print(f"CEPs importados nesta execucao: {imported_ceps}")
        print(f"Produtos importados nesta execucao: {imported_products}")
        print(f"Relacoes produto x CEP importadas nesta execucao: {imported_relations}")
        print(f"Totais atuais -> CEPs: {total_ceps} | Produtos: {total_products} | Relacoes: {total_rel}")
    finally:
        conn.close()


if __name__ == "__main__":
    main()
