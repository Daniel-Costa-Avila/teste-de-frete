"""Gera um arquivo de importacao (Grupo;Nome do produto;ID do produto;Link do produto;CEPs para testar)
a partir da amostra reduzida de representantes por faixa de peso, cruzando com os links
ja cadastrados em 'Base de teste.xlsx'.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]


def load_links(base_path: Path) -> dict[str, str]:
    wb = load_workbook(base_path, read_only=True, data_only=True)
    ws = wb.active
    header = [str(c.value or "").strip().lower() for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx_id = header.index("id_produto")
    idx_link = header.index("link")
    links: dict[str, str] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        pid = str(row[idx_id] or "").strip()
        link = str(row[idx_link] or "").strip()
        if pid and link:
            links[pid] = link
    return links


def main() -> int:
    json_path = ROOT / "Amostra Frete por Faixa de Peso.json"
    base_path = ROOT / "Base de teste.xlsx"
    out_path = ROOT / "artifacts" / "amostra_faixa_peso_importacao.xlsx"
    out_path.parent.mkdir(parents=True, exist_ok=True)

    data = json.loads(json_path.read_text(encoding="utf-8"))
    representantes = data["Representantes"]
    links = load_links(base_path)

    wb = Workbook()
    ws = wb.active
    ws.title = "Produtos"
    ws.append(["Grupo", "Nome do produto", "ID do produto", "Link do produto", "CEPs para testar"])

    missing: list[str] = []
    written = 0
    for item in representantes:
        pid = str(item.get("SKU") or "").strip()
        if not pid or pid == "—":
            continue
        link = links.get(pid, "")
        if not link:
            missing.append(pid)
            continue
        grupo = str(item.get("Faixa de peso") or "").strip()
        nome = str(item.get("Nome do Produto") or "").strip()
        ws.append([grupo, nome, pid, link, ""])
        written += 1

    wb.save(out_path)

    print(f"Arquivo gerado: {out_path}")
    print(f"Produtos escritos: {written}")
    if missing:
        print(f"SKUs sem link na Base de teste.xlsx ({len(missing)}): {', '.join(missing)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
